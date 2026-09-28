from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


MAX_TEXT_BYTES = 32 * 1024
MAX_XLSX_CELLS = 4000


def collect_project_context(root: Path) -> dict[str, object]:
    """Collect read-only project context with provenance for AI review.

    The collector intentionally favors explicit project-local sources. It does
    not fetch external documents or infer rules that are not present in files.
    """

    root = root.resolve()
    requirements = _collect_text_documents(
        root,
        _unique_paths(
            [
                root / "requirements",
                root / "docs" / "requirements.md",
                root / "README.md",
            ],
            suffixes={".md", ".rst", ".txt"},
        ),
    )
    project_rules = _collect_text_documents(
        root,
        _unique_paths(
            [
                root / "project-rules",
                root / "AGENTS.md",
                root / ".github" / "copilot-instructions.md",
            ],
            suffixes={".md", ".rst", ".txt"},
        ),
    )

    design_items: list[dict[str, object]] = []
    warnings: list[str] = []
    for path in _unique_paths(
        [
            root / "design",
            root / "docs" / "design",
            root / "docs" / "architecture.md",
        ],
        suffixes={".xlsx", ".csv", ".md", ".txt", ".pdf", ".xls"},
    ):
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            parsed, parse_warnings = _parse_xlsx(root, path)
            design_items.extend(parsed)
            warnings.extend(parse_warnings)
        elif suffix in {".csv", ".md", ".txt"}:
            item = _read_text_document(root, path)
            if item:
                item["document_type"] = suffix.removeprefix(".")
                design_items.append(item)
        elif suffix in {".xls", ".pdf"}:
            warnings.append(
                f"{_relative(root, path)}: 構造化解析には対応していません。OCR/テキストを提供するか、.xlsxへ変換してください。"
            )

    return {
        "requirements": requirements,
        "design_context": design_items,
        "project_rules": project_rules,
        "warnings": warnings,
    }


def _unique_paths(
    roots: Iterable[Path],
    *,
    suffixes: set[str],
) -> list[Path]:
    results: list[Path] = []
    seen: set[Path] = set()

    for candidate in roots:
        if candidate.is_file():
            paths = [candidate]
        elif candidate.is_dir():
            paths = sorted(path for path in candidate.rglob("*") if path.is_file())
        else:
            continue

        for path in paths:
            if path.suffix.lower() not in suffixes:
                continue
            resolved = path.resolve()
            if resolved in seen:
                continue
            seen.add(resolved)
            results.append(resolved)
    return results


def _collect_text_documents(root: Path, paths: list[Path]) -> list[dict[str, object]]:
    results: list[dict[str, object]] = []
    for path in paths:
        item = _read_text_document(root, path)
        if item:
            results.append(item)
    return results


def _read_text_document(root: Path, path: Path) -> dict[str, object] | None:
    data = path.read_bytes()
    truncated = len(data) > MAX_TEXT_BYTES
    data = data[:MAX_TEXT_BYTES]

    text: str | None = None
    encoding = "utf-8"
    for candidate in ("utf-8", "cp932"):
        try:
            text = data.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = data.decode("utf-8", errors="replace")
        encoding = "utf-8-replacement"

    return {
        "source_file": _relative(root, path),
        "content": text,
        "encoding": encoding,
        "truncated": truncated,
    }


def _parse_xlsx(root: Path, path: Path) -> tuple[list[dict[str, object]], list[str]]:
    warnings: list[str] = []
    items: list[dict[str, object]] = []
    workbook = load_workbook(path, read_only=True, data_only=False)
    consumed = 0

    try:
        for worksheet in workbook.worksheets:
            sheet_cells: list[dict[str, object]] = []
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    sheet_cells.append(
                        {
                            "cell": cell.coordinate,
                            "value": cell.value,
                            "data_type": cell.data_type,
                        }
                    )
                    consumed += 1
                    if consumed >= MAX_XLSX_CELLS:
                        warnings.append(
                            f"{_relative(root, path)}: Excelの抽出件数が上限（{MAX_XLSX_CELLS}件の非空セル）に達しました。"
                        )
                        break
                if consumed >= MAX_XLSX_CELLS:
                    break

            if sheet_cells:
                items.append(
                    {
                        "source_file": _relative(root, path),
                        "document_type": "xlsx",
                        "sheet": worksheet.title,
                        "sheet_state": worksheet.sheet_state,
                        "cells": sheet_cells,
                        "provenance": {
                            "source_file": _relative(root, path),
                            "sheet": worksheet.title,
                        },
                        "limitations": [
                            "図形・画像の意味は構造化しない",
                            "色・罫線など視覚表現の意味は判定しない",
                        ],
                    }
                )
            if consumed >= MAX_XLSX_CELLS:
                break
    finally:
        workbook.close()

    return items, warnings


def _relative(root: Path, path: Path) -> str:
    return path.resolve().relative_to(root).as_posix()
